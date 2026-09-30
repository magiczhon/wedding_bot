"""Сценарные тесты бота: запуск — `pytest` из корня проекта"""
import sqlite3
from io import BytesIO

import openpyxl

from core import bot as b
from core.event_config import default_config
from database.db import Database


# ---------- согласие и удаление данных ----------

def test_consent_is_asked_before_name_and_decline_stores_nothing(env):
    ev = env.event()
    env.open_code('vk', 100, ev['code'])
    consent = env.vk.last(100)
    assert consent[3] == [[b.BTN_AGREE], [b.BTN_DISAGREE]]
    assert 'ИП Тестов Т.Т.' in consent[2] and ev['title'] in consent[2]
    assert 'https://svc.example/privacy' in consent[2] and '/delete' in consent[2]

    env.say('vk', 100, 'Петр Иванов')  # без согласия имя не принимается
    assert env.vk.last(100)[3] == [[b.BTN_AGREE], [b.BTN_DISAGREE]]
    env.say('vk', 100, b.BTN_DISAGREE)
    assert env.db.guests(ev['id']) == [] and env.db.users_by_platform_ids('vk', ['100']) == []
    assert [c['action'] for c in env.db.consent_log()] == ['declined']

    env.guest('vk', 100, 'Петр Иванов', ev)
    user = env.db.users_by_platform_ids('vk', ['100'])[0]
    assert user['full_name'] == 'Петр Иванов' and user['consent_version'] == '1'
    assert env.db.consent_log()[0]['action'] == 'given'


def test_delete_command_removes_everything_and_logs_withdrawal(env):
    ev = env.event()
    env.team(ev)
    env.guest('telegram', 100, 'Петр Иванов', ev, rsvp='yes')
    env.say('telegram', 100, b.BTN_PREFS)
    env.say('telegram', 100, b.FOR_ME_PREFIX + 'Петр Иванов')
    env.say('telegram', 100, '🐟 Рыба')
    env.say('telegram', 100, '🥂 Вино белое')
    env.say('telegram', 100, ev['config']['guest_channels']['host']['button'])
    env.say('telegram', 100, 'Вопрос ведущему')

    env.say('telegram', 100, b.BTN_PRIVACY)
    assert 'Петр Иванов' in env.tg.texts(100)[-1] and ev['title'] in env.tg.texts(100)[-1]
    env.say('telegram', 100, b.BTN_DELETE_ME)
    env.say('telegram', 100, b.BTN_CANCEL)
    assert len(env.db.guests(ev['id'])) == 1

    env.say('telegram', 100, '/delete')
    env.say('telegram', 100, b.BTN_DELETE_YES)
    assert env.db.guests(ev['id']) == [] and env.db.preferences(ev['id']) == []
    assert not any('Вопрос ведущему' in (m['text'] or '') for m in env.db.messages(ev['id']))
    assert env.db.consent_log()[0]['action'] == 'withdrawn'
    env.say('telegram', 100, 'привет')  # бот снова просит ссылку или код
    assert 'код мероприятия' in env.tg.texts(100)[-1]


def test_new_consent_version_is_asked_again(env):
    ev = env.event()
    env.guest('telegram', 100, 'Петр Иванов', ev)
    env.bot.privacy['version'] = '2'
    env.say('telegram', 100, b.BTN_DRESS)
    assert env.tg.last(100)[3] == [[b.BTN_AGREE], [b.BTN_DISAGREE]]
    env.say('telegram', 100, b.BTN_AGREE)
    assert env.db.users_by_platform_ids('telegram', ['100'])[0]['consent_version'] == '2'
    # гость ещё не ответил на приглашение — после согласия бот продолжает с вопроса о присутствии
    assert env.tg.last(100)[3] == [[b.RSVP_YES, b.RSVP_NO], [b.RSVP_MAYBE], [b.BTN_MENU]]


# ---------- гости ----------

def test_guest_registration_rsvp_and_preferences_without_duplicates(env):
    ev = env.event()
    env.open_code('vk', 100, ev['code'])
    env.say('vk', 100, b.BTN_AGREE)
    assert 'имя и фамилию' in env.vk.texts(100)[-1]
    env.say('vk', 100, 'Петр')
    assert 'И фамилию' in env.vk.texts(100)[-1]
    env.say('vk', 100, 'Петр Иванов')
    assert env.vk.last(100)[3] == [[b.RSVP_YES, b.RSVP_NO], [b.RSVP_MAYBE], [b.BTN_MENU]]
    env.say('vk', 100, b.RSVP_YES)
    env.say('vk', 100, '3')
    guest = env.db.guests(ev['id'], 'yes')[0]
    assert (guest['full_name'], guest['party_size'], guest['platform']) == ('Петр Иванов', 3, 'vk')

    for name, food, drink in [(b.FOR_ME_PREFIX + 'Петр Иванов', '🐟 Рыба', '🥂 Вино белое'),
                              ('Ольга Иванова', '🥩 Мясо', b.BTN_SKIP),
                              (b.FOR_ME_PREFIX + 'Петр Иванов', '🥩 Мясо', '🥃 Коньяк')]:
        env.say('vk', 100, b.BTN_PREFS)
        env.say('vk', 100, name)
        env.say('vk', 100, food)
        env.say('vk', 100, drink)
    prefs = {p['guest_name']: (p['food'], p['drink']) for p in env.db.preferences(ev['id'])}
    assert prefs == {'Петр Иванов': ('🥩 Мясо', '🥃 Коньяк'), 'Ольга Иванова': ('🥩 Мясо', 'не указано')}


def test_unknown_user_is_asked_for_code_and_wrong_code_is_rejected(env):
    env.say('telegram', 100, '/start')
    assert 'код мероприятия' in env.tg.texts(100)[-1]
    env.say('telegram', 100, 'код ZZZZZZ')
    assert 'не найден' in env.tg.texts(100)[-1]


def test_guest_does_not_see_staff_actions(env):
    ev = env.event()
    env.guest('telegram', 100, 'Петр Иванов', ev)
    env.say('telegram', 100, b.BTN_EXPORT)
    assert env.tg.texts(100)[-1].startswith('Не понял')
    env.say('telegram', 100, b.BTN_PLACE)
    assert env.tg.last(100)[0] == 'location'


def test_guest_limit_and_archived_event(env):
    ev = env.event(guest_limit=1)
    env.guest('telegram', 100, 'Петр Иванов', ev)
    env.open_code('telegram', 101, ev['code'])
    assert 'лимит гостей' in env.tg.texts(101)[-1]
    assert len(env.db.guests(ev['id'])) == 1

    env.db.update_event(ev['id'], status='archived')
    env.say('telegram', 100, b.BTN_DRESS)
    assert 'завершено' in env.tg.texts(100)[-1]
    env.open_code('telegram', 102, ev['code'])
    assert 'завершено' in env.tg.texts(102)[-1]


# ---------- несколько мероприятий ----------

def test_events_are_isolated(env):
    a, c = env.event('Свадьба А'), env.event('Свадьба Б')
    env.staff('telegram', 1, 'Мария Организатор', a, 'organizer')
    env.staff('telegram', 2, 'Павел Организатор', c, 'organizer')
    env.guest('telegram', 100, 'Гость Первый', a)
    env.guest('vk', 200, 'Гость Второй', c)

    env.say('telegram', 1, b.BTN_BROADCAST)
    env.say('telegram', 1, 'Всем гостям')
    env.say('telegram', 1, 'Сбор в 15:30')
    assert '(1 чел.)' in env.tg.texts(1)[-1]
    env.say('telegram', 1, b.BTN_CONFIRM)
    assert 'Сбор в 15:30' in env.tg.texts(100)[-1]
    assert not any('15:30' in t for t in env.vk.texts(200))

    env.say('telegram', 1, b.BTN_EXPORT)
    wb = openpyxl.load_workbook(BytesIO(env.tg.last(1)[2][1]))
    names = [r[1] for r in wb['Гости'].iter_rows(min_row=2, values_only=True) if r[0]]
    assert names == ['Гость Первый']
    # организатор Б не видит гостей А в сводке
    env.say('telegram', 2, b.BTN_SUMMARY)
    assert 'Зарегистрировано гостей: 1' in env.tg.texts(2)[-1] and 'Свадьба Б' in env.tg.texts(2)[-1]


def test_host_on_two_weddings_gets_tagged_questions_and_replies_route_by_event(env):
    a, c = env.event('Свадьба А'), env.event('Свадьба Б')
    env.staff('vk', 20, 'Олег Ведущий', a, 'host')
    env.staff('vk', 20, 'Олег Ведущий', c, 'host')  # сейчас у ведущего открыто мероприятие Б
    assert env.db.get_user(env.db.users_by_platform_ids('vk', ['20'])[0]['id'])['current_event_id'] == c['id']
    assert [b.BTN_PRIVACY, b.BTN_EVENTS] == env.vk.last(20)[3][-1]

    env.guest('telegram', 100, 'Петр Иванов', a)
    env.say('telegram', 100, a['config']['guest_channels']['host']['button'])
    env.say('telegram', 100, 'Можно песню для мамы?')
    question = env.vk.last(20)
    assert question[2].startswith('[Свадьба А]') and 'песню для мамы' in question[2]

    env.say('vk', 20, 'Конечно!', reply_to=question[4])  # ответ уходит в свадьбу А
    assert 'Конечно!' in env.tg.texts(100)[-1] and 'Олег Ведущий' in env.tg.texts(100)[-1]
    guest_id = env.db.guests(a['id'])[0]['id']
    env.say('vk', 20, f'ответ {guest_id} Жду!')
    assert 'Жду!' in env.tg.texts(100)[-1]
    assert [m['text'] for m in env.db.messages(c['id'])] == []  # в журнал Б ничего не попало

    env.say('vk', 20, b.BTN_EVENTS)
    env.say('vk', 20, 'Свадьба А')
    assert 'Свадьба А' in env.vk.texts(20)[-1]


def test_question_to_host_across_messengers_and_answers(env):
    ev = env.event()
    env.team(ev)
    env.guest('telegram', 100, 'Петр Иванов', ev)
    guest_id = env.db.guests(ev['id'])[0]['id']

    env.say('telegram', 100, ev['config']['guest_channels']['host']['button'])
    env.say('telegram', 100, 'Можно заказать песню для мамы?')
    assert env.tg.texts(100)[-1].startswith('✅ Передано')
    question = env.vk.last(20)
    assert f'гостя №{guest_id}' in question[2] and not question[2].startswith('[')
    assert not any('песню' in t for t in env.tg.texts(1))

    env.say('vk', 20, 'Конечно, после торта!', reply_to=question[4])
    answer = env.tg.last(100)
    assert 'после торта' in answer[2]
    env.say('telegram', 100, b.BTN_MENU)
    env.say('telegram', 100, 'Спасибо!', reply_to=answer[4])
    assert 'Спасибо!' in env.vk.texts(20)[-1]


def test_host_broadcast_needs_confirmation_and_guest_reply_goes_to_host(env):
    ev = env.event()
    env.team(ev)
    env.guest('telegram', 100, 'Петр Иванов', ev)
    env.guest('vk', 101, 'Ирина Смирнова', ev)

    env.say('vk', 20, b.BTN_WRITE_GUESTS)
    env.say('vk', 20, 'Кто хочет участвовать в конкурсе?')
    assert 'всем гостям (2 чел.)' in env.vk.texts(20)[-1]
    env.say('vk', 20, b.BTN_CANCEL)
    assert not any('конкурсе' in t for t in env.tg.texts(100))
    env.say('vk', 20, 'Кто хочет участвовать в конкурсе?')
    env.say('vk', 20, b.BTN_CONFIRM)
    assert 'Отправлено: 2' in env.vk.texts(20)[-1]
    broadcast = env.tg.last(100)
    env.say('telegram', 100, 'Я хочу!', reply_to=broadcast[4])
    assert 'Я хочу!' in env.vk.texts(20)[-1]


def test_contractor_room_reaches_only_its_members_with_media(env):
    from core.messenger import Attachment
    ev = env.event()
    env.team(ev)
    env.say('telegram', 40, b.BTN_ROOMS)
    assert [row[0] for row in env.tg.last(40)[3]] == ['Общий чат команды', 'Декор', b.BTN_MENU]
    env.say('telegram', 40, 'Декор')
    photo = Attachment('photo', load=lambda: b'PNG', filename='arka.jpg')
    env.say('telegram', 40, 'Вариант арки', attachments=[photo])
    for uid in (1, 30):
        assert [e[0] for e in env.tg.to(uid)[-2:]] == ['text', 'copy']
    assert not any('арки' in t for t in env.vk.texts(20))
    env.say('telegram', 30, 'Нравится!', reply_to=env.tg.to(30)[-1][4])
    assert 'Нравится!' in env.tg.texts(40)[-1]

    env.say('telegram', 40, b.BTN_MENU)
    env.say('telegram', 40, b.BTN_ROOMS)
    env.say('telegram', 40, 'Общий чат команды')
    env.say('telegram', 40, 'Схема зала', attachments=[photo])
    assert [e[0] for e in env.vk.to(20)[-2:]] == ['text', 'photo']


def test_organizer_tools_summary_export_and_invites(env):
    ev = env.event()
    env.team(ev)
    env.guest('telegram', 100, 'Петр Иванов', ev, rsvp='yes')
    env.guest('vk', 101, 'Ирина Смирнова', ev)
    env.guest('telegram', 102, 'Сергей Козлов', ev, rsvp='no')

    env.say('telegram', 1, b.BTN_BROADCAST)
    env.say('telegram', 1, 'Кто не ответил')
    env.say('telegram', 1, 'Напоминаем ответить до пятницы')
    env.say('telegram', 1, b.BTN_CONFIRM)
    assert 'до пятницы' in env.vk.texts(101)[-1]
    assert not any('до пятницы' in t for t in env.tg.texts(100) + env.tg.texts(102))
    env.say('vk', 101, 'Приду!', reply_to=env.vk.last(101)[4])
    assert 'Приду!' in env.tg.texts(1)[-1] and 'Приду!' in env.tg.texts(30)[-1]

    env.say('telegram', 1, b.BTN_SUMMARY)
    assert 'Придут: 1 (всего человек: 2)' in env.tg.texts(1)[-1]
    env.say('telegram', 1, b.BTN_INVITES)
    invites = env.tg.texts(1)[-1]
    assert f"telegram-guest:{ev['code']}" in invites and 'Ведущий — код' in invites

    env.say('telegram', 30, b.BTN_EXPORT)
    wb = openpyxl.load_workbook(BytesIO(env.tg.last(30)[2][1]))
    assert wb.sheetnames == ['Гости', 'Пожелания', 'Команда', 'Сообщения']


def test_staff_guest_view_and_back(env):
    ev = env.event()
    env.team(ev)
    env.say('telegram', 30, b.BTN_GUEST_VIEW)
    assert [b.BTN_RSVP] in env.tg.last(30)[3]
    env.say('telegram', 30, b.BTN_DRESS)
    assert 'Дресс-код' in env.tg.texts(30)[-1]
    env.say('telegram', 30, b.BTN_STAFF_VIEW)
    assert env.tg.last(30)[3][0][0] == b.BTN_ROOMS


def test_web_broadcast_reaches_guests_and_replies_go_to_organizers(env):
    ev = env.event()
    env.team(ev)
    env.guest('telegram', 100, 'Петр Иванов', ev)
    account = {'id': 1}
    sent, failed = env.bot.web_broadcast(ev['id'], 'Всем гостям', 'Трансфер в 15:00', account)
    assert (sent, failed) == (1, 0) and 'Трансфер в 15:00' in env.tg.texts(100)[-1]
    env.say('telegram', 100, 'Где остановка?', reply_to=env.tg.last(100)[4])
    assert 'Где остановка?' in env.tg.texts(1)[-1]


# ---------- перенос данных старых версий ----------

def test_import_of_first_version_database(tmp_path):
    path = str(tmp_path / 'old.db')
    conn = sqlite3.connect(path)
    conn.executescript('''
        CREATE TABLE guests (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER UNIQUE, chat_id INTEGER,
                             username TEXT, tg_name TEXT, full_name TEXT, updated_at TIMESTAMP);
        CREATE TABLE guest_preferences (id INTEGER PRIMARY KEY AUTOINCREMENT, guest_id INTEGER, guest_name TEXT,
                                        food TEXT, drink TEXT, updated_at TIMESTAMP);
        INSERT INTO guests (user_id, chat_id, username, tg_name, full_name) VALUES (555, 555, 'petr', 'Петя', 'Петр Иванов');
        INSERT INTO guests (user_id, chat_id, username, tg_name, full_name) VALUES (777, 777, 'boss', 'Босс', 'Мария Орг');
        INSERT INTO guest_preferences (guest_id, guest_name, food, drink) VALUES (1, 'Петр Иванов', 'Рыба', 'Водка');
        INSERT INTO guest_preferences (guest_id, guest_name, food, drink) VALUES (1, 'Петр Иванов', 'Мясо', 'Водка');
    ''')
    conn.commit()
    conn.close()
    db = Database(path)
    assert db.has_legacy_data()
    event = db.import_legacy(default_config(), admins={'telegram': [777]})
    assert [g['full_name'] for g in db.guests(event['id'])] == ['Петр Иванов']
    assert [m['full_name'] for m in db.members(event['id'], roles=['organizer'])] == ['Мария Орг']
    assert [(p['guest_name'], p['food']) for p in db.preferences(event['id'])] == [('Петр Иванов', 'Мясо')]
    assert not db.has_legacy_data() and db.import_legacy(default_config()) is None


def test_import_of_single_event_multiplatform_database(tmp_path):
    path = str(tmp_path / 'v1.db')
    conn = sqlite3.connect(path)
    conn.executescript('''
        CREATE TABLE users (id INTEGER PRIMARY KEY AUTOINCREMENT, platform TEXT, platform_user_id TEXT, chat_id TEXT,
            username TEXT, display_name TEXT, full_name TEXT, role TEXT, rsvp TEXT, party_size INTEGER, state TEXT,
            state_data TEXT, created_at TIMESTAMP, updated_at TIMESTAMP);
        CREATE TABLE preferences (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER, guest_name TEXT, food TEXT,
            drink TEXT, updated_at TIMESTAMP);
        CREATE TABLE invites (role TEXT PRIMARY KEY, code TEXT);
        CREATE TABLE relays (id INTEGER PRIMARY KEY, platform TEXT, chat_id TEXT, message_id TEXT, channel TEXT,
            target_user_id INTEGER, recipient_user_id INTEGER, created_at TIMESTAMP);
        CREATE INDEX relays_lookup ON relays (platform, chat_id, message_id);
        CREATE TABLE messages (id INTEGER PRIMARY KEY, channel TEXT, sender_user_id INTEGER, target_user_id INTEGER,
            text TEXT, created_at TIMESTAMP);
        INSERT INTO users (platform, platform_user_id, chat_id, full_name, role, rsvp, party_size)
            VALUES ('vk', '20', '20', 'Олег Ведущий', 'host', NULL, NULL),
                   ('telegram', '100', '100', 'Петр Иванов', 'guest', 'yes', 2);
        INSERT INTO preferences (user_id, guest_name, food, drink) VALUES (2, 'Петр Иванов', 'Рыба', 'Вино');
        INSERT INTO invites VALUES ('host', 'HOST42');
        INSERT INTO messages (channel, sender_user_id, target_user_id, text) VALUES ('guest:host', 2, NULL, 'Песню!');
    ''')
    conn.commit()
    conn.close()
    db = Database(path)
    event = db.import_legacy(default_config())
    guest = db.guests(event['id'])[0]
    assert (guest['full_name'], guest['rsvp'], guest['party_size']) == ('Петр Иванов', 'yes', 2)
    assert db.members(event['id'], roles=['host'])[0]['full_name'] == 'Олег Ведущий'
    assert db.resolve_code('HOST42')[2] == 'host'
    assert db.messages(event['id'])[0]['text'] == 'Песню!'
    Database(path)  # повторный запуск ничего не ломает
    assert len(db.guests(event['id'])) == 1
