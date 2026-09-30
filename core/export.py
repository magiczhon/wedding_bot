"""Выгрузка в Excel: гости с ответами, пожелания по меню, команда и журнал сообщений"""
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

RSVP_TITLES = {'yes': 'Придёт', 'no': 'Не придёт', 'maybe': 'Пока не знает', None: 'Не ответил'}
RSVP_FILLS = {'yes': 'E2F0D9', 'no': 'F8D7DA', 'maybe': 'FFF3CD'}


def _sheet(wb, title, headers, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in rows:
        ws.append(list(row))
    for column in ws.columns:
        width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        ws.column_dimensions[column[0].column_letter].width = min(width + 2, 60)
    ws.freeze_panes = 'A2'
    return ws


def build_excel(db, bot, event) -> bytes:
    """Выгрузка одного мероприятия. bot — WeddingBot: названия мессенджеров и ссылки на профили"""
    config = event['config']
    roles = config.get('roles', {})
    channels = config.get('guest_channels') or {}
    rooms = {r['id']: r for r in config.get('rooms') or []}

    def role_title(role):
        return 'Гость' if role == 'guest' else roles.get(role, role or '')

    def link(user):
        m = bot.messengers.get(user['platform'])
        return m.profile_link(user) if m else ''

    wb = openpyxl.Workbook()

    guests = db.guests(event['id'])
    ws = _sheet(wb, 'Гости', ['№', 'ФИО', 'Ответ', 'Сколько человек', 'Мессенджер', 'Профиль', 'Обновлено'],
                [(g['id'], g['full_name'], RSVP_TITLES.get(g['rsvp'], g['rsvp']),
                  g['party_size'] if g['rsvp'] == 'yes' else '', bot.platform_title(g['platform']), link(g),
                  str(g['updated_at'])[:16]) for g in guests], first=True)
    for row, g in zip(ws.iter_rows(min_row=2), guests):
        if g['rsvp'] in RSVP_FILLS:
            row[2].fill = PatternFill('solid', fgColor=RSVP_FILLS[g['rsvp']])
    total = sum(g['party_size'] or 1 for g in guests if g['rsvp'] == 'yes')
    ws.append([])
    ws.append(['', 'Всего придут (человек)', '', total])
    ws.cell(ws.max_row, 2).font = Font(bold=True)

    _sheet(wb, 'Пожелания', ['Гость', 'Еда', 'Напиток', 'Кто заполнил', 'Мессенджер', 'Обновлено'],
           [(p['guest_name'], p['food'], p['drink'], p['owner_name'], bot.platform_title(p['platform']),
             str(p['updated_at'])[:16]) for p in db.preferences(event['id'])])

    team = db.members(event['id'], exclude_guests=True)
    _sheet(wb, 'Команда', ['Роль', 'ФИО', 'Мессенджер', 'Профиль'],
           [(role_title(t['role']), t['full_name'], bot.platform_title(t['platform']), link(t)) for t in team])

    def channel_title(channel):
        kind, _, cid = channel.partition(':')
        if kind == 'guest':
            return channels.get(cid, {}).get('title', cid) + ' (гости)'
        if kind == 'room':
            return 'Чат: ' + rooms.get(cid, {}).get('title', cid)
        return 'Рассылка: ' + cid

    _sheet(wb, 'Сообщения', ['Время', 'Канал', 'От кого', 'Роль', 'Кому', 'Текст'],
           [(str(m['created_at'])[:16], channel_title(m['channel']), m['sender_name'] or '',
             role_title(m['sender_role']) if m['sender_role'] else 'Веб-интерфейс',
             m['target_name'] or '', m['text']) for m in db.messages(event['id'])])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
